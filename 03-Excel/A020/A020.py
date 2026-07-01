import openpyxl
myBook=openpyxl.load_workbook('利润表.xlsx')
mySheet=myBook['1季度利润表']
#循环1季度利润表的第5行到44行
for i in range(5,45):
    #将每行的第1列的“(万元)”修改为“(元)”
    myValue=mySheet.cell(i,1).value
    myValue=myValue[0:-4]+'(元)'
    mySheet.cell(i,1).value=myValue
    #将每行的第2列的金额乘以10000
    myValue=mySheet.cell(i,2).value
    if(myValue is not None):
          myValue=myValue*10000
          mySheet.cell(i,2).value=myValue
myBook.save('结果表-利润表.xlsx')
