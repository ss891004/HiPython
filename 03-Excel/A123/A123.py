import openpyxl
myBook=openpyxl.load_workbook('收入表.xlsx')
mySheet=myBook.active
mySheet.unmerge_cells('A2:A4')
mySheet['A3']=mySheet['A2'].value
mySheet['A4']=mySheet['A2'].value
mySheet.unmerge_cells('A5:A7')
mySheet['A6']=mySheet['A5'].value
mySheet['A7']=mySheet['A5'].value
mySheet.unmerge_cells('A8:A10')
mySheet['A9']=mySheet['A8'].value
mySheet['A10']=mySheet['A8'].value
mySheet.unmerge_cells('A11:A13')
mySheet['A12']=mySheet['A11'].value
mySheet['A13']=mySheet['A11'].value
myBook.save('结果表-收入表.xlsx')
