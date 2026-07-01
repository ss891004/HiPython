import openpyxl
#myBook=openpyxl.load_workbook('收入表.xlsx')
myBook=openpyxl.load_workbook('收入表.xlsx',data_only=True)
mySheet=myBook.active
myRows=mySheet.rows
#新建工作簿(myNewBook)，在保存之后即为新的Excel文件
myNewBook=openpyxl.Workbook()
myNewSheet=myNewBook.active
myNewSheet.title='收入表'
#循环收入表的行(myRow)
for myRow in myRows:
    myList=[]
    #循环行(myRow)的单元格
    for myCell in myRow:
        myList+=[myCell.value]
    #在新工作簿的收入表(myNewSheet)中添加行(数据)
    myNewSheet.append(myList)
myNewBook.save('结果表-收入表.xlsx')
