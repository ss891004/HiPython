import openpyxl
myBook=openpyxl.load_workbook('成绩表.xlsx')
mySheet=myBook['成绩表']
#数据范围(myRange)从成绩表(mySheet)的第2行开始，到最后一行
myRange=mySheet[str(mySheet.min_row+1):str(mySheet.max_row)]
#新建工作表(差等生表)
myFilterSheet=myBook.create_sheet('差等生表')
#在差等生表中添加表头
myFilterSheet.append([myCell.value for myCell
    in mySheet[str(mySheet.min_row):str(mySheet.min_row)]]+['总分'])
for myRow in myRange:
    #获取每位学生的各科成绩
    myList=[myCell.value for myCell in myRow]
    #计算每位学生的成绩总分
    myScore=sum(myList[1:])
    #如果总分小于350，则添加到差等生表
    if myScore<350:
       myFilterSheet.append(myList+[myScore])
myBook.save('结果表-成绩表.xlsx')
