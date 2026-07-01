import openpyxl
import random
myBook=openpyxl.load_workbook('录取表.xlsx',data_only=True)
mySheet=myBook.active
myNewBook=openpyxl.Workbook()
myNewSheet=myNewBook.active
myNewSheet.title='录取表'
#在新工作表(录取表)中添加第1列(年份列)，该列不参与随机排列
myRowIndex=1
for myCell in list(mySheet.columns)[0]:
    myNewSheet.cell(myRowIndex,1).value=myCell.value;
    myRowIndex+=1
#获取录取表(mySheet)的所有列(第1列除外)
myCols=list(mySheet.columns)[1:]
#随机排列录取表的所有列(第1列除外)
random.shuffle(myCols)
#输出随机排列结果
myColIndex=2
for myCol in myCols:
    myRowIndex=1
    for myCell in myCol:
        myNewSheet.cell(myRowIndex,myColIndex).value=myCell.value;
        myRowIndex+=1
    myColIndex+=1
myNewBook.save('结果表-录取表.xlsx')
