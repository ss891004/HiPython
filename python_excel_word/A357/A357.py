import openpyxl
myBook=openpyxl.load_workbook('收入表.xlsx')
mySheet=myBook.worksheets[0]
#获取收入表(mySheet)的第5行到第8行的所有单元格数据(myList)
myList=list(mySheet.values)[4:8]
#获取第1列(家电收入列)的B5:B8范围的所有单元格数据(myListC1)
myListC1=[myList[i][1] for i in range(len(myList))]
#对第1列(家电收入列)的单元格数据求和，并将合计写入最后一行的单元格
mySheet.cell(9,2).value=sum(myListC1)
#获取第2列(建材收入列)的C5:C8范围的所有单元格数据(myListC2)
myListC2=[myList[i][2] for i in range(len(myList))]
mySheet.cell(9,3).value=sum(myListC2)
#获取第3列(其他收入列)的D5:D8范围的所有单元格数据(myListC3)
myListC3=[myList[i][3] for i in range(len(myList))]
mySheet.cell(9,4).value=sum(myListC3)
myBook.save('结果表-收入表.xlsx')
