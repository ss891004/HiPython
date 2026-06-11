import docx
myDocument=docx.Document('打折商品.docx')
#使用分区样式设置Word文件的第1节的页眉(注意：从第0节开始)
myDocument.sections[1].header.is_linked_to_previous=False
myParagraph1=myDocument.sections[1].header.paragraphs[0]
myParagraph1.text = "人气指数：👨‍👨‍👧\t第1节 进口水果\t推荐指数：❤‍❤‍❤"
myParagraph1.style =myDocument.styles["Header"]
myParagraph1.runs[0].font.size=docx.shared.Pt(14)
#使用分区样式设置Word文件的第2节的页眉
myDocument.sections[2].header.is_linked_to_previous=False
myParagraph2=myDocument.sections[2].header.paragraphs[0]
myParagraph2.text = "人气指数：👨‍👨‍👧👨‍👨‍👧\t第2节 当季蔬菜\t推荐指数：❤‍❤❤‍❤"
myParagraph2.style =myDocument.styles["Header"]
myParagraph2.runs[0].font.size=docx.shared.Pt(14)
#使用分区样式设置Word文件的第3节的页眉
myDocument.sections[3].header.is_linked_to_previous=False
myParagraph3=myDocument.sections[3].header.paragraphs[0]
myParagraph3.text = "人气指数：👨‍👨‍👧\t第3节 川渝火锅\t推荐指数：❤‍❤‍❤"
# myParagraph3.text = "第3节 川渝火锅\t\t推荐指数：❤‍❤‍❤"
myParagraph3.style =myDocument.styles["Header"]
myParagraph3.runs[0].font.size=docx.shared.Pt(14)
myDocument.save('我的Word文件-打折商品.docx')
