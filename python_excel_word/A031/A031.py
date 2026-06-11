import openpyxl
myBook=openpyxl.load_workbook('收入表.xlsx')
mySheet=myBook.active
#指定收入表(mySheet)的数据范围(myRange)，即B5：D8
myRange=mySheet.iter_rows(min_row=5,min_col=2,max_row=8,max_col=4)
myRowIndex=5
#循环数据范围(myRange)的行(myRow)
for myRow in myRange:
    #对行(myRow)的单元格数据求和(myRowSum)
    myRowSum=sum([myCell.value for myCell in myRow])
    mySheet.cell(myRowIndex,5).value=myRowSum
    myRowIndex+=1
myBook.save('结果表-收入表.xlsx')
