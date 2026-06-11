import openpyxl
myBook=openpyxl.load_workbook('收入表.xlsx')
mySheet=myBook['收入表']
myNewBook=openpyxl.Workbook()
myNewSheet=myNewBook.active
myNewSheet.title='收入表'
myNewSheet.append(['季度','业务类别','营业收入'])
myRange1=mySheet['A'][4:]
myRange2=mySheet.iter_rows(min_col=2,min_row=5)
for myQuarter,myRow in list(map(lambda x,y:[x,y],myRange1,myRange2)):
    for myType,myAmount in list(map(lambda x,y:[x,y],mySheet['4'][1:],myRow)):
        myNewSheet.append([myQuarter.value, myType.value, myAmount.value])
myNewBook.save('结果表-收入表.xlsx')
