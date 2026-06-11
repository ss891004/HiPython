import openpyxl
myBook=openpyxl.load_workbook('收入表.xlsx')
mySheet=myBook['收入表']
mySheet.merge_cells("A2:A4")
mySheet.merge_cells("A5:A7")
mySheet.merge_cells("A8:A10")
mySheet.merge_cells("A11:A13")
myBook.save('结果表-收入表.xlsx')
