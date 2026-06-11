import docx
myDocument=docx.Document('打折商品.docx')
#在Word文件中创建从偶数页开始的第1节
myDocument.add_section(docx.enum.section.WD_SECTION.EVEN_PAGE)
myParagraph1=myDocument.add_paragraph("第1节 进口水果")
myDocument.add_picture('image11.jpg',width=docx.shared.Inches(4.0))
myDocument.add_picture('image12.jpg',width=docx.shared.Inches(4.0))
#在Word文件中创建从偶数页开始的第2节
myDocument.add_section(docx.enum.section.WD_SECTION.EVEN_PAGE)
myParagraph2=myDocument.add_paragraph("第2节 当季蔬菜")
myDocument.add_picture('image21.jpg',width=docx.shared.Inches(4.0))
myDocument.add_picture('image22.jpg',width=docx.shared.Inches(4.0))
#在Word文件中创建从偶数页开始的第3节
myDocument.add_section(docx.enum.section.WD_SECTION.EVEN_PAGE)
myParagraph3=myDocument.add_paragraph("第3节 川渝火锅")
myDocument.add_picture('image31.jpg',width=docx.shared.Inches(4.0))
myDocument.add_picture('image32.jpg',width=docx.shared.Inches(4.0))
#设置Word文件的字体
for myParagraph in myDocument.paragraphs:
   myParagraph.paragraph_format.alignment=\
                         docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
   for myRun in  myParagraph.runs:
       myRun.font.name='Times New Roman'
       myRun.font.element.rPr.rFonts.set(docx.oxml.ns.qn('w:eastAsia'),'楷体')
       myRun.font.size=docx.shared.Pt(36)
myDocument.save('我的Word文件-打折商品.docx')
