import docx
myDocument=docx.Document('长寿湖简介.docx')
#获取Word文件(myDocument)的第2个段落的文本(myText)
myText=myDocument.paragraphs[1].text
#在文本(myText)中将'长寿湖'替换为'长寿湖景区'
myText=myText.replace('长寿湖','长寿湖景区')
#使用替换之后的文本 (myText)设置第2个段落
myDocument.paragraphs[1].text=myText
myDocument.save('我的Word文件-长寿湖简介.docx')
