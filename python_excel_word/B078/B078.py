import docx
myDocument=docx.Document('快捷键.docx')
myTable=myDocument.tables[0]
myTable.cell(1,0).width=docx.shared.Inches(2.5)
myDocument.save('我的Word文件-快捷键.docx')
