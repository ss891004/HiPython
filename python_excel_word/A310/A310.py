import openpyxl
#读取“录取表.xlsx”文件
myBook=openpyxl.load_workbook('录取表.xlsx')
mySheet=myBook['录取表']
#按行获取录取表(mySheet)的单元格数据(myRange)
myRange=list(mySheet.values)
#创建空白字典(myDict)
myDict={}
#从录取表(myRange)的第4行开始循环(到最后一行)
for myRow in myRange[3:]:
    #如果在字典(myDict)中存在某录取院校且存在该录取院校的某专业
    if myRow[0] in myDict.keys() and myRow[1] in myDict[myRow[0]].keys():
         #则直接在某录取院校的某专业中添加考生([myRow])
         myDict[myRow[0]][myRow[1]]+=[myRow]
    #否则
    else:
         #如果在字典(myDict)中不存在某录取院校，则首先创建某录取院校及其专业
         if  myRow[0] not in myDict.keys(): myDict[myRow[0]]={}
         myDict[myRow[0]][myRow[1]]=[myRow]
#循环字典(myDict)的成员
for myKey1,myValue1 in myDict.items():
    #创建新工作簿(myNewBook)
    myNewBook=openpyxl.Workbook()
    for myKey2, myValue2 in myValue1.items():
        #根据键名(myKey2)创建新工作表(myNewSheet)
        myNewSheet=myNewBook.create_sheet(myKey2+"专业录取表")
        #在新工作表(myNewSheet)中添加表头(录取院校、专业、考生姓名、总分)
        myNewSheet.append(myRange[2])
        #在新工作表(myNewSheet)中添加键名(专业)下的多个键值(考生)
        for myRow in myValue2:
            myNewSheet.append(myRow)
    myNewBook.remove(myNewBook['Sheet'])
    #保存拆分之后(各个录取院校)的工作簿(myNewBook)，即保存各个Excel文件
    myPath='结果表-'+myKey1+'录取表.xlsx'
    myNewBook.save(myPath)
