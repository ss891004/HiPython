import docx
myDocument=docx.Document('新员工.docx')
myTable=myDocument.tables[0]
myTable.cell(1,0).merge(myTable.cell(2,0)).merge(myTable.cell(3,0))
myTable.cell(1,0).text='投资部'
myDocument.save('我的Word文件-新员工.docx')
