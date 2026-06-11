import docx
myDocument=docx.Document('散文名篇.docx')
#获取Word文件(myDocument)的第2个段落
myParagraph1=myDocument.paragraphs[1]._element
#删除Word文件(myDocument)的第2个段落
myParagraph1.getparent().remove(myParagraph1)
#获取Word文件(myDocument)的第3个段落(即获取原文件的第4个段落)
myParagraph2=myDocument.paragraphs[2]._element
#删除Word文件(myDocument)的第3个段落(即删除原文件的第4个段落)
myParagraph2.getparent().remove(myParagraph2)
myDocument.save('我的Word文件-散文名篇.docx')
