import openpyxl
myBook=openpyxl.load_workbook('成绩表.xlsx')
mySheet=myBook['成绩表']
#数据范围(myRange)从成绩表(mySheet)的第5行开始，到最后一行
myRange=mySheet[str(mySheet.min_row+4):str(mySheet.max_row)]
#新建工作表(汇总表)
mySumSheet=myBook.create_sheet('汇总表')
#在汇总表(mySumSheet)中添加表头
mySumSheet.append(['姓名','总分'])
for myRow in myRange:
    #在汇总表(mySumSheet)中添加每个学生的姓名和总分
    mySumSheet.append([myRow[0].value,
                       sum([myCell.value for myCell in myRow][1:])])
myBook.save('结果表-成绩表.xlsx')
