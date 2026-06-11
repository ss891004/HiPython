import openpyxl
myBook=openpyxl.load_workbook('收入表.xlsx')
mySheet=myBook.active
#自定义青色的左细斜纹线填充样式(myPatternFill)
myPatternFill=openpyxl.styles.PatternFill(fill_type='lightDown',
                                          fgColor='97ffff')
#循环收入表(mySheet)的['A5:D8']范围的行(myRow)
for myRow in mySheet['A5:D8']:
    #循环行(myRow)的单元格(myCell)
    for myCell in myRow:
        #使用自定义填充样式(myPatternFill)设置单元格(myCell)的fill属性
        myCell.fill=myPatternFill
myBook.save('结果表-收入表.xlsx')
