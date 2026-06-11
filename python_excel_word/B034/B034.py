import docx
myDocument=docx.Document('唐诗宋词精品.docx')
#在Word文件(myDocument)中新建第1个标题
myDocument.add_heading(text="第1部分 唐诗精品",level=1)
#在第1个标题下新建段落
myDocument.add_paragraph(text='唐诗，泛指创作于唐朝诗人的诗，为唐代儒客文人之智慧佳作。唐诗是中华民族珍贵的文化遗产之一，是中华文化宝库中的一颗明珠，同时也对世界上许多国家的文化发展产生了很大影响，对于后人研究唐代的政治、民情、风俗、文化等都有重要的参考意义。')
#新建第2个标题
myDocument.add_heading(text="第2部分 宋词精品")
#在第2个标题下新建段落
myDocument.add_paragraph(text='宋词是一种相对于古体诗的新体诗歌之一，为宋代儒客文人智慧精华，标志宋代文学的最高成就。宋词句子有长有短，便于歌唱。因是合乐的歌词，故又称曲子词、乐府、乐章、长短句、诗余、琴趣等。')
myDocument.save('我的Word文件-唐诗宋词精品.docx')
