import docx
#自定义函数设置单元格背景
def fillCellColor(row,col,color):
    myElement='<w:shd {} w:fill="{color_value}"/>'
    myFormat=myElement.format(docx.oxml.ns.nsdecls('w'),color_value=color)
    myXML=docx.oxml.parse_xml(myFormat)
    myTable.cell(row,col)._tc.get_or_add_tcPr().append(myXML)
myDocument=docx.Document('新员工.docx')
#获取Word文件(myDocument)的第1个表格(myTable)
myTable=myDocument.tables[0]
#循环第1个表格(myTable)的每个单元格
for i in range(len(myTable.rows)):
    for j in range(len(myTable.columns)):
        #设置奇数行的单元格背景颜色
        if(i%2==0):
           fillCellColor(i,j,'#F5F5F5')
        #设置偶数行的单元格背景颜色
        else:
           fillCellColor(i,j,'#E0FFFF')
myDocument.save('我的Word文件-新员工.docx')
