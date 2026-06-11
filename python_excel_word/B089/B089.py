import docx
myDocument=docx.Document('销量榜.docx')
myTable=myDocument.tables[0]
myRow=myTable.add_row();
myRow.cells[1].add_paragraph().add_run().add_picture('image1.jpg')
myRow.cells[2].add_paragraph().add_run().add_picture('image2.jpg')
myDocument.save('我的Word文件-销量榜.docx')
