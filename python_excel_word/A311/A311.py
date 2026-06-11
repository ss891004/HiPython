import openpyxl
myBook=openpyxl.load_workbook('收入表.xlsx')
mySheet=myBook['收入表']
myNewBook=openpyxl.Workbook()
myNewSheet=myNewBook.active
myNewSheet.title='收入表'
myNewSheet.append(['季度','业务类别','营业收入'])
myRange1=mySheet['A'][4:]
myRange2=mySheet.iter_rows(min_col=2,min_row=5)
for myQuarter,myRow in zip(myRange1,myRange2):
    #print([myQuarter.value]+[myCell.value for myCell in myRow])
    for myType, myAmount in zip(mySheet['4'][1:],myRow):
        #print(myQuarter.value, myType.value, myAmount.value)
        myNewSheet.append([myQuarter.value, myType.value, myAmount.value])
myNewBook.save('结果表-收入表.xlsx')
