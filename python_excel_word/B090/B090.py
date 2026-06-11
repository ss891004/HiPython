import docx
myDocument=docx.Document('新学员.docx')
#获取Word文件(myDocument)的第1个表格(myTable)
myTable=myDocument.tables[0]
i=len(myTable.rows)-1
#循环第1个表格(myTable)的每行
while(i>0):
   #删除联系地址不包含'渝北区'的学员
   #即剩下的学员则为联系地址包含'渝北区'的学员
   if  '渝北区' not in  myTable.cell(i,2).text:
       myRow=myTable.rows[i]
       myRow._element.getparent().remove(myRow._element)
   i=i-1
myDocument.save('我的Word文件-新学员.docx')
