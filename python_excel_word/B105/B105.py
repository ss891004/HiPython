import docx
myDocument=docx.Document('年度收入.docx')
#获取Word文件(myDocument)的第1个表格(myTable)
myTable=myDocument.tables[0]
#循环第1个表格(myTable)的第2、3、4、5列
for j in range(1,len(myTable.columns)):
    mySum=0
    #循环每列的第2、3、4、5行
    for i in range(1,len(myTable.rows)-1):
        #累加每列的第2、3、4、5行的单元格数据
        mySum+=int(myTable.column_cells(j)[i].text)
        #在每列的第6行的单元格中写入各个类别的收入合计
        myTable.column_cells(j)[5].text=str(mySum)
        #myTable.cell(5,j).text=str(mySum)
myDocument.save('我的Word文件-年度收入.docx')
