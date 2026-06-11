import openpyxl
myBook=openpyxl.load_workbook('收入表.xlsx')
mySheet=myBook.active
#根据起始单元格和结束单元格设置数据范围的最小行列数和最大行列数
myMinRow=mySheet.min_row+4
myMinColumn=mySheet.min_column+1
myMaxRow=mySheet.max_row
myMaxColumn=mySheet.max_column
#指定按列操作单元格的数据范围
myRange=mySheet.iter_cols(min_row=myMinRow,min_col=myMinColumn,
                          max_row=myMaxRow,max_col=myMaxColumn)
myRowIndex=myMinRow+4
myColIndex=myMinColumn-1
mySheet.cell(myRowIndex,myColIndex).value='合计'
#按列对单元格数据求和
for myCol in myRange:
    myColIndex+=1
    myColSum=sum(myCell.value for myCell in myCol)
    #将合计数据写入单元格
    mySheet.cell(myRowIndex,myColIndex).value=myColSum
myBook.save('结果表-收入表.xlsx')
