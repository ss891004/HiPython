import docx
myDocument=docx.Document('散文名篇.docx')
#设置将要添加的文本(myText)
myText="再让我们回顾一下故事的开头，不正是乔的新邻居告诉他地里出黄金的吗?然而，事实上是乔对英国的语言理解得还不够透彻。他的新邻居其实是说他那块土地有肥沃的土壤，所以你应该知道黄金的概念来自哪儿了吧。"
#在Word文件(myDocument)的末尾添加文本(追加新的段落)
myParagraph=myDocument.add_paragraph(myText,style='Body Text')
# myParagraph=myDocument.add_paragraph(myText,style='Body Text 3')
# for s in myDocument.styles:
#     print(s.name)
myDocument.save('我的Word文件-散文名篇.docx')
