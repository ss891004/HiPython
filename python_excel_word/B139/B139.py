import docx
myDocument=docx.Document('唐宋名篇.docx')
myDocument.add_paragraph('将进酒','Title')
myDocument.add_paragraph('君不见，黄河之水天上来，奔流到海不复回。君不见，高堂明镜悲白发，朝如青丝暮成雪。')
myDocument.save('我的Word文件-唐宋名篇.docx')
