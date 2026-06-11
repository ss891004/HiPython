import docx
myDocument=docx.Document('快捷键.docx')
#获取Word文件(myDocument)的第1个表格(myTable)
myTable=myDocument.tables[0]
#循环第1个表格(myTable)的所有行
for i in range(len(myTable.rows)):
     #设置每行第1列的单元格的文本右对齐
     myTable.cell(i,0).paragraphs[0].alignment=\
                   docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
     # # 设置每行第1列的单元格的文本左对齐
     # myTable.cell(i,0).paragraphs[0].alignment=\
     #  docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT
myDocument.save('我的Word文件-快捷键.docx')
