import docx
myDocument=docx.Document('散文名篇.docx')
myParagraph2=myDocument.paragraphs[1].insert_paragraph_before()
myImage=myParagraph2.add_run().add_picture('myimage.jpg')
myImage.height=docx.shared.Cm(1)
myImage.width=docx.shared.Cm(10.7)
myDocument.save('我的Word文件-散文名篇.docx')
