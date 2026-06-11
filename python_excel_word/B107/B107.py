import docx
myDocument=docx.Document()
#print(len(myDocument.sections))
myDocument.add_paragraph("打折商品推荐书")
#新建第1节(注意：默认将自动创建一个节，即第0节)
mySection1=myDocument.add_section()
myParagraph1=myDocument.add_paragraph("第1节 进口水果")
# myParagraph1.add_run().add_picture('image11.jpg')
myDocument.add_picture('image11.jpg')
myDocument.add_picture('image12.jpg')
#新建第2节
myDocument.add_section()
myParagraph2=myDocument.add_paragraph("第2节 当季蔬菜")
# myParagraph2.add_run().add_picture('image21.jpg')
myDocument.add_picture('image21.jpg')
myDocument.add_picture('image22.jpg')
#新建第3节
myDocument.add_section()
myParagraph3=myDocument.add_paragraph("第3节 川渝火锅")
# myParagraph3.add_run().add_picture('image31.jpg')
myDocument.add_picture('image31.jpg')
myDocument.add_picture('image32.jpg')
#设置Word文件的字体
for myParagraph in myDocument.paragraphs:
    myParagraph.paragraph_format.alignment=\
                          docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    for myRun in  myParagraph.runs:
       myRun.font.name='Times New Roman'
       myRun.font.element.rPr.rFonts.set(docx.oxml.ns.qn('w:eastAsia'),'楷体')
       myRun.font.size=docx.shared.Pt(36)
myDocument.save('我的Word文件-打折商品.docx')
