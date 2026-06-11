import docx
myDocument=docx.Document('虞美人.docx')
myDocument.paragraphs[1].runs[0].font.strike=True
myDocument.paragraphs[2].runs[0].font.double_strike=True
myDocument.save('我的Word文件-虞美人.docx')
