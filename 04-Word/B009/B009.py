import docx
myDocument=docx.Document('散文名篇.docx')
myDocument.paragraphs[1].paragraph_format.space_before=docx.shared.Inches(0.5)
myDocument.paragraphs[1].paragraph_format.space_after=docx.shared.Inches(0.5)
myDocument.save('我的Word文件-散文名篇.docx')
