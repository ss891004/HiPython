import docx
myDocument=docx.Document('凤梨简介.docx')
# #在Word文件(myDocument)的末尾添加新段落(myParagraph2)
# myParagraph2=myDocument.add_paragraph()
# #在新段落(myParagraph2)中添加指定的图像(myImage)
# myImage=myParagraph2.add_run().add_picture('myimage.png')
#在Word文件(myDocument)的末尾添加新段落，并在新段落中添加图像(myImage)
myImage=myDocument.add_picture('myimage.png')
#自定义图像(myImage)的高度和宽度
myImage.height=docx.shared.Cm(4)
myImage.width=docx.shared.Cm(3)
#居中对齐图像(即居中对齐第2个段落)
myDocument.paragraphs[1].alignment=\
                 docx.enum.text.WD_PARAGRAPH_ALIGNMENT.CENTER
# #右对齐图像(即右对齐第2个段落)
# myDocument.paragraphs[1].alignment=\
#                  docx.enum.text.WD_PARAGRAPH_ALIGNMENT.RIGHT
myDocument.save('我的Word文件-凤梨简介.docx')
