import docx
myDocument=docx.Document('散文名篇.docx')
#myDocument.paragraphs[0].paragraph_format.left_indent=docx.shared.Inches(2)
myDocument.paragraphs[0].paragraph_format.right_indent=docx.shared.Inches(2)
myDocument.save('我的Word文件-散文名篇.docx')
