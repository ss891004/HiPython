import docx
myDocument=docx.Document('唐诗宋词精品.docx')
#在Word文件(myDocument)中新建一级标题(myHeading)
myHeading=myDocument.add_heading(level=1)
#设置一级标题(myHeading)居中对齐
myHeading.alignment=docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
# #设置一级标题(myHeading)右对齐
# myHeading.alignment=docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
#添加一级标题(myHeading)的文本
myRun=myHeading.add_run("第1部分 唐诗精品")
#设置一级标题(myHeading)的字体大小
myRun.font.size=docx.shared.Pt(12)
#设置一级标题(myHeading)的中文字体
myRun.font.name='微软雅黑'
myRun.element.rPr.rFonts.set(docx.oxml.ns.qn('w:eastAsia'),'微软雅黑')
#初始化6个列表项
myItems=["九月九日忆山东兄弟","梦游天姥吟留别","凉州词",
         "闻官军收河南河北","白雪歌送武判官归京","春江花月夜"]
for i in range(6):
    #(在一级标题(myHeading)下面)新建有序列表
    myParagraph=myDocument.add_paragraph(text=myItems[i], style='List Number')
myDocument.save('我的Word文件-唐诗宋词精品.docx')
