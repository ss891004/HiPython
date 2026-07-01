import docx
myDocument=docx.Document('新员工.docx')
#获取Word文件(myDocument)的第1个表格(myTable)
myTable=myDocument.tables[0]
#循环第1个表格(myTable)的每个单元格
for i in range(len(myTable.rows)):
    for j in range(len(myTable.columns)):
        #设置单元格的文本颜色为蓝色
        myTable.cell(i,j).paragraphs[0].runs[0].font.color.rgb=\
                                docx.shared.RGBColor(55,55,255)
        # #设置单元格的文本颜色为红色
        # myTable.cell(i,j).paragraphs[0].runs[0].font.color.rgb=\
        #  docx.shared.RGBColor(255,55,55)
myDocument.save('我的Word文件-新员工.docx')
