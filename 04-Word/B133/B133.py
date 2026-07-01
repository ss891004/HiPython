import docx
myDocument=docx.Document('散文名篇.docx')
#print(myDocument.paragraphs[1].text)
#实际效果是两段文本颜色均为红色
myDocument.paragraphs[1].style.font.color.rgb=docx.shared.RGBColor(255,55,55)
#实际效果是两段文本颜色均为绿色
#myDocument.paragraphs[0].style.font.color.rgb=docx.shared.RGBColor(0,255,0)
#实际效果是只有第2个段落的文本颜色为红色
#myDocument.paragraphs[1].runs[0].font.color.rgb=docx.shared.RGBColor(255,0,0)
myDocument.save('我的Word文件-散文名篇.docx')
