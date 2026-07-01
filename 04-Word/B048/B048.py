import docx
myDocument=docx.Document('散文名篇.docx')
#在第2个段落(myDocument.paragraphs[1])的前面插入新的段落(myParagraph2)
myParagraph2=myDocument.paragraphs[1].insert_paragraph_before("    Did Joe's sons become slaves to the digging? No, they were inspired because they had visions of what money could do for them and did not focus on the money itself. ")
#在Word文件(myDocument)的末尾添加段落(myParagraph4)
myParagraph4=myDocument.add_paragraph("    乔的儿子们成为挖掘土地的奴隶了吗?不，他们只是一昧地幻想着金钱能为他们做什么，并被无数个幻想所激励，但从未考虑过获取金钱的正确途径。")
myDocument.save('我的Word文件-散文名篇.docx')
