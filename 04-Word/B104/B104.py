import docx
myDocument=docx.Document('年度收入.docx')
#获取Word文件(myDocument)的第1个表格(myTable)
myTable=myDocument.tables[0]
#循环第1个表格(myTable)的第2、3、4、5行
for i in range(1,len(myTable.rows)):
    mySum=0
    #循环每行的第2、3、4列
    for j in range(1,len(myTable.columns)-1):
      #累加每行的第2、3、4列的单元格数据
      mySum+=int(myTable.row_cells(i)[j].text)
      #在每行的第5列的单元格中写入合计
      myTable.row_cells(i)[4].text=str(mySum)
      #myTable.cell(i,4).text=str(mySum)
myDocument.save('我的Word文件-年度收入.docx')
