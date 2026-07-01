import docx
myDocument=docx.Document('表情符号.docx')
myTable=myDocument.tables[0]
# myTable.cell(1,1).vertical_alignment=docx.enum.table.WD_ALIGN_VERTICAL.TOP
# myTable.cell(2,1).vertical_alignment=docx.enum.table.WD_ALIGN_VERTICAL.TOP
myTable.cell(1,1).vertical_alignment=docx.enum.table.WD_ALIGN_VERTICAL.CENTER
myTable.cell(2,1).vertical_alignment=docx.enum.table.WD_ALIGN_VERTICAL.CENTER
myDocument.save('我的Word文件-表情符号.docx')
