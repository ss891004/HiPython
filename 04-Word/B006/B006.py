import docx
myDocument=docx.Document('散文名篇.docx')
#获取Word文件(myDocument)的第2个段落的文本(myText)
myText=myDocument.paragraphs[1].text
#在文本(myText)中设置每个单词的首字母大写，其余小写
myText=myText.title()
# #在文本(myText)中设置每个单词的所有字母大写和小写互换
# myText=myText.swapcase()
# #在文本(myText)中设置每个单词的所有字母小写
# myText=myText.lower()
# #在文本(myText)中设置每个单词的所有字母大写
# myText=myText.upper()
#使用替换之后的文本(myText)重新设置第2个段落的text属性
myDocument.paragraphs[1].text=myText
myDocument.save('我的Word文件-散文名篇.docx')
