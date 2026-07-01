import docx
myDocument=docx.Document('春节祝福.docx')
i=0
#循环Word文件(myDocument)的第1个段落的块(myRun)
for myRun in myDocument.paragraphs[0].runs:
    if i<11:
        #在块(myRun)中添加图像
        myImage=myRun.add_picture('myimage'+str(i)+'.png')
    i+=1
myDocument.save('我的Word文件-春节祝福.docx')
