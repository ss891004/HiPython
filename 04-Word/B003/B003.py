import docx
myDocument=docx.Document('散文名篇.docx')
myText=myDocument.paragraphs[1].text
myDocument.paragraphs[1].text=myDocument.paragraphs[2].text
myDocument.paragraphs[2].text=myText
myDocument.save('我的Word文件-散文名篇.docx')
