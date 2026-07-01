import docx
#创建自定义单元格边框颜色的函数
def setCellBorderColor(row,col,myColor):
    myTcPr=myTable.cell(row,col)._tc.get_or_add_tcPr()
    myTcBorders=myTcPr.first_child_found_in("w:tcBorders")
    if myTcBorders is None:
       myTcBorders=docx.oxml.OxmlElement('w:tcBorders')
       myTcPr.append(myTcBorders)
    for myEdge in ('left','top','right','bottom'):
        myEdgeData={"color":myColor}
        if myEdgeData:
           myTag='w:{}'.format(myEdge)
           myElement=myTcBorders.find(docx.oxml.ns.qn(myTag))
           if myElement is None:
              myElement=docx.oxml.OxmlElement(myTag)
           myTcBorders.append(myElement)
           myElement.set(docx.oxml.ns.qn('w:{}'.format("color")),
                                str(myEdgeData["color"]))
myDocument=docx.Document('新员工.docx')
#获取Word文件(myDocument)的第1个表格(myTable)
myTable=myDocument.tables[0]
#将第1个表格(myTable)的所有单元格的边框线设置为红色
for i in range(len(myTable.rows)):
    for j in range(len(myTable.columns)):
        setCellBorderColor(i,j,"#ff0000")
myDocument.save('我的Word文件-新员工.docx')
