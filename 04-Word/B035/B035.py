import docx
myDocument=docx.Document('唐诗宋词精品.docx')
#在Word文件(myDocument)中新建第1个一级标题
myDocument.add_heading(text="第1部分 唐诗精品",level=1)
#在Word文件(myDocument)中新建第1个一级标题下面的五级标题101
myDocument.add_heading(text="101 九月九日忆山东兄弟",level=5)
#在Word文件(myDocument)中新建第1个一级标题下面的五级标题102
myDocument.add_heading(text="102 梦游天姥吟留别",level=5)
#在Word文件(myDocument)中新建第1个一级标题下面的五级标题103
myDocument.add_heading(text="103 凉州词",level=5)
#在Word文件(myDocument)中新建第1个一级标题下面的五级标题104
myDocument.add_heading(text="104 闻官军收河南河北",level=5)
#在Word文件(myDocument)中新建第1个一级标题下面的五级标题105
myDocument.add_heading(text="105 白雪歌送武判官归京",level=5)
#在Word文件(myDocument)中新建第1个一级标题下面的五级标题106
myDocument.add_heading(text="106 春江花月夜",level=5)
#在Word文件(myDocument)的第1个一级标题下面的五级标题106中新建段落
myDocument.add_paragraph(text='春江潮水连海平，海上明月共潮生。滟滟随波千万里，何处春江无月明！江流宛转绕芳甸，月照花林皆似霰。空里流霜不觉飞，汀上白沙看不见。江天一色无纤尘，皎皎空中孤月轮。')
#在Word文件(myDocument)中新建第2个一级标题
myDocument.add_heading(text="第2部分 宋词精品")
myDocument.save('我的Word文件-唐诗宋词精品.docx')
