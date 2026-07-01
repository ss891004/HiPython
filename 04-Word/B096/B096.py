import docx
#创建自定义单元格边框的函数
def setCellBorder(row,col,**borderArgs):
    myTcPr=myTable.cell(row,col)._tc.get_or_add_tcPr()
    myTcBorders=myTcPr.first_child_found_in("w:tcBorders")
    if myTcBorders is None:
       myTcBorders=docx.oxml.OxmlElement('w:tcBorders')
       myTcPr.append(myTcBorders)
    for myEdge in ('left','top','right','bottom'):
        myEdgeData=borderArgs.get(myEdge)
        if myEdgeData:
           myTag='w:{}'.format(myEdge)
           myElement=myTcBorders.find(docx.oxml.ns.qn(myTag))
           if myElement is None:
              myElement=docx.oxml.OxmlElement(myTag)
           myTcBorders.append(myElement)
           for myKey in ["sz","val","color"]:
               if myKey in myEdgeData:
                  myElement.set(docx.oxml.ns.qn('w:{}'.format(myKey)),
                                str(myEdgeData[myKey]))
myDocument=docx.Document('新员工.docx')
#获取Word文件(myDocument)的第1个表格(myTable)
myTable=myDocument.tables[0]
#循环第1个表格(myTable)的每个单元格
for i in range(len(myTable.rows)):
    for j in range(len(myTable.columns)):
        #自定义单元格的边框线：
        #sz表示边框粗细、val表示边框类型、color表示边框颜色
        setCellBorder(i,j,top={"sz":12,"val":"double","color":"#ff0000"},
                          bottom={"sz":12,"val":"double","color":"#ff0000"},
                          left={"sz":12,"val":"double","color":"#ff0000"},
                          right={"sz":12,"val":"double","color":"#ff0000"})
myDocument.save('我的Word文件-新员工.docx')
