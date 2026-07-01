import docx
myDocument=docx.Document('散文名篇.docx')
#清空Word文件(myDocument)的第2个段落的文本
myDocument.paragraphs[1].text=''
#在第2个段落(myDocument.paragraphs[1])中添加1个块
myDocument.paragraphs[1].add_run('    再让我们回顾一下故事的开头，不正是乔的新邻居告诉他地里出黄金的吗?')
#在第2个段落(myDocument.paragraphs[1])中再添加1个块(myRun)
myRun=myDocument.paragraphs[1].add_run('然而，事实上是乔对英国的语言理解得还不够透彻。')
#设置块(myRun)的高亮颜色为黄色
myRun.font.highlight_color=docx.enum.text.WD_COLOR_INDEX.YELLOW
# #设置块(myRun)的高亮颜色为红色
# myRun.font.highlight_color=docx.enum.text.WD_COLOR_INDEX.RED
#在第2个段落(myDocument.paragraphs[1])中再添加1个块
myDocument.paragraphs[1].add_run('他的新邻居其实是说他那块土地有肥沃的土壤，所以你应该知道黄金的概念来自哪儿了吧。')
myDocument.save('我的Word文件-散文名篇.docx')
