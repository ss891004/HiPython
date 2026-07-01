import docx
myDocument=docx.Document('唐诗宋词精品.docx')
#在Word文件(myDocument)中新建第1个一级标题
myDocument.add_heading(text="第1部分 唐诗精品",level=1)
#在Word文件(myDocument)中新建2个无序列表项，且左端缩进1个字符
myDocument.add_paragraph(text="九月九日忆山东兄弟",style='List Bullet')
myDocument.add_paragraph(text="梦游天姥吟留别",style='List Bullet')
#在Word文件(myDocument)中新建2个无序列表项，且左端缩进2个字符
myDocument.add_paragraph(text="凉州词",style='List Bullet 2')
myDocument.add_paragraph(text="闻官军收河南河北",style='List Bullet 2')
#在Word文件(myDocument)中新建2个无序列表项，且左端缩进3个字符
myDocument.add_paragraph(text="白雪歌送武判官归京",style='List Bullet 3')
myDocument.add_paragraph(text="春江花月夜",style='List Bullet 3')
#在Word文件(myDocument)中新建第2个一级标题
myDocument.add_heading(text="第2部分 宋词精品")
myDocument.save('我的Word文件-唐诗宋词精品.docx')
